import io
import unicodedata
from decimal import Decimal, InvalidOperation

import openpyxl
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.models.medicamento import Medicamento
from app.repositories.lote_repository import LoteRepository
from app.repositories.medicamento_repository import MedicamentoRepository
from app.schemas.medicamento import (
    MedicamentoCreate,
    MedicamentoImportarValorAtualizado,
    MedicamentoImportarValorLinhaInvalida,
    MedicamentoImportarValorResultado,
    MedicamentoUpdate,
)


def _normalizar_nome(texto: str) -> str:
    """minúsculo, sem acento, espaços simples — mesmo critério usado na
    importação de valor unitário do Almoxarifado (ItemService), pra
    comparar nome de planilha contra nome do catálogo sem depender de
    digitação exata."""
    texto = " ".join(str(texto).strip().split()).lower()
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


class MedicamentoService:
    """Cadastro de medicamentos — Farmacêutico cadastra sozinho, sem
    aprovação do Coordenador (docs/00_PROJETO.md seção 9, default 5)."""

    def __init__(self):
        self.repository = MedicamentoRepository()
        self.lote_repository = LoteRepository()

    def create(self, db: Session, dados: MedicamentoCreate):
        return self.repository.create(db, dados)

    def list(self, db: Session, apenas_ativos: bool = True):
        return self.repository.list(db, apenas_ativos)

    def get_by_id(self, db: Session, medicamento_id: int):
        medicamento = self.repository.get_by_id(db, medicamento_id)

        if medicamento is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Medicamento não encontrado.",
            )

        return medicamento

    def update(self, db: Session, medicamento_id: int, dados: MedicamentoUpdate):
        medicamento = self.get_by_id(db, medicamento_id)

        return self.repository.update(db, medicamento, dados)

    def gerar_template_valor_unitario(self, db: Session) -> io.BytesIO:
        """Planilha modelo para a importação de valor unitário em massa
        (2026-09-18, pedido do cliente) — só Nome + Valor Unitário Novo:
        sem coluna de "valor atual" porque não existe um valor único do
        medicamento (o preço mora em cada lote, que varia por compra)."""
        medicamentos = self.repository.list(db, apenas_ativos=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Valor unitário"
        ws.append(["Nome", "Valor Unitário Novo"])
        for medicamento in medicamentos:
            ws.append([medicamento.nome, None])
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 22

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def importar_valor_unitario(self, db: Session, conteudo: bytes) -> MedicamentoImportarValorResultado:
        """Lê a planilha preenchida e aplica o valor unitário por NOME do
        medicamento (sem diferenciar maiúsculo/acento) em TODOS os lotes
        ATIVOS (com saldo) daquele medicamento, em qualquer unidade —
        pedido do cliente: a planilha vale como "preço atual desse
        medicamento". Cada linha é resolvida de forma independente: uma
        linha problemática não derruba a importação inteira, só entra no
        relatório final pra revisão manual."""
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
        except Exception as exc:  # noqa: BLE001 — qualquer falha de parsing vira 400 com mensagem clara
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Não foi possível ler o arquivo como planilha .xlsx: {exc}",
            ) from exc

        ws = wb.active
        linhas = [[c.value for c in row] for row in ws.iter_rows()]
        if not linhas:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Planilha vazia.")

        cabecalho = [_normalizar_nome(c) if c is not None else "" for c in linhas[0]]
        try:
            idx_nome = cabecalho.index("nome")
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Coluna 'Nome' não encontrada na primeira linha da planilha.",
            ) from None
        idx_valor = next(
            (i for i, c in enumerate(cabecalho) if c in ("valor unitario novo", "valor unitario", "valor")),
            None,
        )
        if idx_valor is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Coluna de valor ('Valor Unitário Novo' ou 'Valor Unitário') não encontrada na planilha.",
            )

        medicamentos_ativos = self.repository.list(db, apenas_ativos=True)
        por_nome: dict[str, list[Medicamento]] = {}
        for medicamento in medicamentos_ativos:
            por_nome.setdefault(_normalizar_nome(medicamento.nome), []).append(medicamento)

        atualizados: list[MedicamentoImportarValorAtualizado] = []
        nao_encontrados: list[str] = []
        ambiguos: list[str] = []
        sem_lote_ativo: list[str] = []
        invalidas: list[MedicamentoImportarValorLinhaInvalida] = []

        for numero_linha, linha in enumerate(linhas[1:], start=2):
            if all(v is None or str(v).strip() == "" for v in linha):
                continue  # linha em branco

            nome_bruto = linha[idx_nome] if idx_nome < len(linha) else None
            if nome_bruto is None or not str(nome_bruto).strip():
                invalidas.append(MedicamentoImportarValorLinhaInvalida(linha=numero_linha, motivo="Nome em branco."))
                continue

            valor_bruto = linha[idx_valor] if idx_valor < len(linha) else None
            if valor_bruto is None or str(valor_bruto).strip() == "":
                continue  # sem valor novo preenchido — linha ignorada silenciosamente

            try:
                valor_novo = Decimal(str(valor_bruto).strip().replace(",", "."))
            except InvalidOperation:
                invalidas.append(
                    MedicamentoImportarValorLinhaInvalida(
                        linha=numero_linha, motivo=f"Valor '{valor_bruto}' não é um número válido."
                    )
                )
                continue
            if valor_novo < 0:
                invalidas.append(
                    MedicamentoImportarValorLinhaInvalida(linha=numero_linha, motivo="Valor não pode ser negativo.")
                )
                continue

            nome = str(nome_bruto).strip()
            candidatos = por_nome.get(_normalizar_nome(nome), [])
            if len(candidatos) == 0:
                nao_encontrados.append(nome)
                continue
            if len(candidatos) > 1:
                ambiguos.append(nome)
                continue

            medicamento = candidatos[0]
            lotes_ativos = self.lote_repository.listar(
                db, medicamento_id=medicamento.id, apenas_disponivel=True, ordenar_fefo=False
            )
            if not lotes_ativos:
                sem_lote_ativo.append(nome)
                continue

            for lote in lotes_ativos:
                lote.valor_unitario = valor_novo
            atualizados.append(
                MedicamentoImportarValorAtualizado(
                    nome=medicamento.nome, lotes_atualizados=len(lotes_ativos), valor_novo=str(valor_novo)
                )
            )

        db.commit()

        return MedicamentoImportarValorResultado(
            atualizados=atualizados,
            nao_encontrados=nao_encontrados,
            ambiguos=ambiguos,
            sem_lote_ativo=sem_lote_ativo,
            invalidas=invalidas,
        )
